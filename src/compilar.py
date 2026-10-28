"""Compila ambos documentos y exporta PowerPoint fiel al PDF final."""
import argparse
from io import BytesIO
from pathlib import Path
import shutil
import subprocess
import pymupdf
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
TARGETS = [ROOT / "articulo/main.tex", ROOT / "presentacion/presentacion-clusters-departamentales.tex"]


def compilar(path):
    engine = shutil.which("pdflatex")
    if not engine:
        raise RuntimeError("Instala una distribución LaTeX con pdflatex (MiKTeX o TeX Live).")
    for _ in range(3):
        result = subprocess.run([engine, "-interaction=nonstopmode", "-halt-on-error", "-file-line-error", path.name], cwd=path.parent, capture_output=True)
        if result.returncode:
            raise RuntimeError((result.stdout + result.stderr).decode("utf-8", errors="replace")[-7000:])
    log = path.with_suffix(".log").read_text(encoding="utf-8", errors="replace")
    problems = [line for line in log.splitlines() if "Overfull" in line or "undefined" in line.lower() or "Missing character:" in line]
    if problems:
        raise RuntimeError(f"Revisar {path}:\n" + "\n".join(problems))
    print(f"Compilación correcta: {path.relative_to(ROOT)}", flush=True)


def powerpoint(pdf):
    doc = pymupdf.open(pdf)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.author = "Círculo de Ciencia de Datos y Econometría (CCDE)"
    presentation.core_properties.title = "Estructura productiva y recuperación regional"
    for index, page in enumerate(doc):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        image = page.get_pixmap(matrix=pymupdf.Matrix(2.5, 2.5), alpha=False)
        picture = slide.shapes.add_picture(BytesIO(image.tobytes("png")), 0, 0, width=presentation.slide_width, height=presentation.slide_height)
        picture._element.nvPicPr.cNvPr.set("descr", f"Diapositiva {index+1}. " + page.get_text().replace("\n", " ")[:2000])
        slide.notes_slide.notes_text_frame.text = page.get_text()
    target = pdf.with_suffix(".pptx")
    presentation.save(target)
    print(f"PowerPoint actualizado: {len(doc)} diapositivas. Fuente editable: LaTeX.", flush=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--solo", choices=["articulo", "presentacion"])
    args = parser.parse_args()
    for path in TARGETS:
        if not args.solo or path.parent.name == args.solo:
            compilar(path)
    if not args.solo or args.solo == "presentacion":
        powerpoint(TARGETS[1].with_suffix(".pdf"))


if __name__ == "__main__":
    main()
