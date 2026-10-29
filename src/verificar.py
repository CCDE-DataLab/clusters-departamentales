"""Audita fuentes, invariantes numéricos y consistencia de entregables finales."""
import argparse
import hashlib
import json
import re
import numpy as np
import pandas as pd
from scipy.spatial.distance import pdist
from analisis import ROOT, cargar_config, ORDER
from composicional import reemplazar_ceros, clr, ilr
from consolidar import extraer, comprobar


def verificar_trazabilidad(config, manifest):
    if config != manifest["config"]:
        raise ValueError("Configuración modificada tras el análisis; regenera las salidas")
    source = ROOT / config["datos"]["entrada"]
    if hashlib.sha256(source.read_bytes()).hexdigest() != manifest["sha256_entrada"]:
        raise ValueError("Entrada modificada tras el análisis")
    sources = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (ROOT / "data/raw").glob("*.xlsx")}
    if sources != manifest["sha256_fuentes"]:
        raise ValueError("El conjunto de fuentes INEI cambió tras el análisis")
    fingerprints = manifest.get("sha256_salidas")
    if not fingerprints:
        raise ValueError("Faltan huellas de las salidas; regenera el análisis")
    output = ROOT / config["datos"]["salida"]
    for name, fingerprint in fingerprints.items():
        path = output / name
        if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != fingerprint:
            raise ValueError(f"Salida modificada o ausente tras el análisis: {name}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", help="Ruta al TOML utilizado para generar el análisis")
    config = cargar_config(parser.parse_args().config)
    output = ROOT / config["datos"]["salida"]
    manifest = json.loads((output / "resultados.json").read_text(encoding="utf-8"))
    source = ROOT / config["datos"]["entrada"]
    verificar_trazabilidad(config, manifest)
    comprobar(extraer(), source)
    panel = pd.read_excel(output / "panel_econometrico_final.xlsx")
    assert not panel.duplicated(["departamento", "anio"]).any()
    c, m = config["clustering"], config["muestra"]
    index_base = panel[panel.anio.eq(m["base"])].VAB_idx2019
    np.testing.assert_allclose(index_base, 100)
    for _, rows in panel.groupby("departamento"):
        rows = rows.sort_values("anio")
        mask = rows.anio.diff().eq(1)
        np.testing.assert_allclose(rows.VAB_total.pct_change(fill_method=None).mul(100).round(2)[mask], rows.crecimiento_vab[mask], atol=.0051)
    shares = pd.read_excel(output / "composiciones_pre.xlsx", index_col=0)
    xclr = pd.read_excel(output / "coordenadas_clr.xlsx", index_col=0)
    xilr = pd.read_excel(output / "coordenadas_ilr.xlsx", index_col=0)
    np.testing.assert_allclose(shares.sum(axis=1), 100)
    np.testing.assert_allclose(xclr.sum(axis=1), 0, atol=1e-12)
    np.testing.assert_allclose(pdist(xclr), pdist(xilr), atol=1e-12)
    # Caso con ceros reales: cierre, positividad e invariancia de cocientes positivos.
    example = pd.DataFrame([[0, 20, 80], [30, 0, 70]])
    replacement = reemplazar_ceros(example, c["delta_ceros"])
    assert replacement.gt(0).all().all()
    np.testing.assert_allclose(replacement.sum(axis=1), 100)
    np.testing.assert_allclose(replacement.iloc[0, 1] / replacement.iloc[0, 2], .25)
    np.testing.assert_allclose(pdist(clr(replacement)), pdist(ilr(replacement)))
    labels = pd.read_excel(output / "etiquetas_clusters.xlsx", index_col=0)
    assert labels.groupby("Cluster_K4").Cluster_Ward_K2.nunique().eq(1).all()
    assert set(labels.Nombre_Cluster_K4) == set(ORDER)
    results = pd.read_excel(output / "resumen_modelos_econometricos.xlsx")
    assert set(results.Modelo) == {f"M{i}" for i in range(1, 7)}
    for code, rows in results.groupby("Modelo"):
        start, end = (m["tasas_inicio"], m["tasas_fin"]) if code in ["M1", "M2", "M3", "M4"] else (m["niveles_inicio"], m["niveles_fin"])
        n = panel[panel.anio.between(start, end)].shape[0] - (1 if code == "M4" else 0)
        assert rows.N_obs.eq(n).all()
        np.testing.assert_allclose(rows.Coeficiente / rows.Error_Estandar, rows.T_stat)
        assert (rows.IC_inferior < rows.Coeficiente).all() and (rows.Coeficiente < rows.IC_superior).all()
        # El contraste de cero y la pertenencia al IC deben coincidir.
        alpha = 1 - config["econometria"]["confianza"]
        exclude_zero = (rows.IC_inferior > 0) | (rows.IC_superior < 0)
        assert np.array_equal(exclude_zero, rows.P_valor < alpha)
    # M4 cambia solamente las interacciones del perfil que contiene la celda excluida.
    one = results[results.Modelo.eq("M1")].set_index("Parametro")
    four = results[results.Modelo.eq("M4")].set_index("Parametro")
    excluded_profile = labels.loc[config["econometria"]["exclusion_departamento"], "Nombre_Cluster_K4"]
    unchanged = [p for p in one.index if p != "Intercept" and f"[T.{excluded_profile}]" not in p]
    if excluded_profile == ORDER[0]:
        unchanged = []  # Al excluir en la referencia también cambian los contrastes relativos.
    np.testing.assert_allclose(one.loc[unchanged, "Coeficiente"], four.loc[unchanged, "Coeficiente"], atol=1e-10)
    macros = (ROOT / "resultados/cifras.tex").read_text(encoding="utf-8")
    values = dict(re.findall(r"\\newcommand\{\\(\w+)\}\{([^{}]+)\}", macros))
    scar = results[(results.Modelo.eq("M5")) & results.Parametro.eq("POST_21_23:Nombre_Cluster_K4[T.Minero/Enclave]")].iloc[0]
    assert values["ScarBeta"] == f"{scar.Coeficiente:.2f}" and values["RatesObs"] == str(manifest["diagnostico"]["n_tasas"])
    for parent in [ROOT / "articulo", ROOT / "presentacion"]:
        for path in parent.rglob("*.tex"):
            text = path.read_text(encoding="utf-8-sig")
            for relative in re.findall(r"\\input\{([^}]+)\}", text):
                assert (parent / relative).exists(), f"Referencia inexistente en {path}: {relative}"
    import pymupdf
    from pptx import Presentation
    presentation = ROOT / "presentacion/presentacion-clusters-departamentales.pdf"
    pdf, paper = pymupdf.open(presentation), pymupdf.open(ROOT / "articulo/main.pdf")
    pptx = Presentation(presentation.with_suffix(".pptx"))
    assert len(pdf) == len(pptx.slides)
    for page, slide in zip(pdf, pptx.slides):
        expected = page.get_pixmap(matrix=pymupdf.Matrix(2.5, 2.5), alpha=False).tobytes("png")
        assert hashlib.sha256(slide.shapes[0].image.blob).digest() == hashlib.sha256(expected).digest(), "PowerPoint desactualizado frente al PDF"
    for doc in [pdf, paper]:
        for page in doc:
            assert "\ufffd" not in page.get_text()
            for word in page.get_text("words"):
                x0, y0, x1, y1 = word[:4]
                assert -1 <= x0 < x1 <= page.rect.width + 1, f"Texto fuera de página: {word}"
                assert -1 <= y0 < y1 <= page.rect.height + 1, f"Texto fuera de página: {word}"
    print(f"Auditoría correcta: fuentes verificadas, {len(results)} coeficientes, {len(pdf)} diapositivas y {len(paper)} páginas de artículo.")


if __name__ == "__main__":
    main()
